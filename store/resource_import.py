from io import BytesIO
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from .models import DigitalResource, Product


HEADERS = [
    'upload row', 'book name', 'book code', 'isbn to paste',
    'student resource link', 'teacher resource link',
]
LEGACY_HEADERS = ['book name', 'isbn', 'student resource link', 'teacher resource link']


def normalize_isbn(value):
    normalized = str(value or '').strip().replace('-', '').replace(' ', '').upper()
    valid_10 = len(normalized) == 10 and normalized[:9].isdigit() and (normalized[-1].isdigit() or normalized[-1] == 'X')
    valid_13 = len(normalized) == 13 and normalized.isdigit()
    return normalized if valid_10 or valid_13 else ''


def template_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Book Reference'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
    sheet.column_dimensions['A'].width = 14
    sheet.column_dimensions['B'].width = 42
    sheet.column_dimensions['C'].width = 16
    sheet.column_dimensions['D'].width = 22
    sheet.column_dimensions['E'].width = 60
    sheet.column_dimensions['F'].width = 60
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:F1'
    instructions = workbook.create_sheet('Instructions')
    instructions.append(['Column', 'How to fill'])
    instructions.append(['upload row', 'Reference only. It is not used to map the book.'])
    instructions.append(['book name', 'Reference only. It helps staff identify the book.'])
    instructions.append(['book code', 'Reference only. It is not used to map the book.'])
    instructions.append(['isbn to paste', 'Required. Must already exist in the books catalogue. Format this cell as Text.'])
    instructions.append(['student resource link', 'Optional public HTTPS destination for students.'])
    instructions.append(['teacher resource link', 'Optional HTTPS destination for teachers.'])
    instructions.append(['Important', 'At least one resource link is required in each row. Uploading the same ISBN updates its links.'])
    for cell in instructions[1]:
        cell.font = Font(bold=True)
    instructions.column_dimensions['A'].width = 30
    instructions.column_dimensions['B'].width = 95
    return workbook


def template_bytes():
    output = BytesIO()
    template_workbook().save(output)
    return output.getvalue()


def _validate_url(value, field, row_number):
    url = str(value or '').strip()
    if not url:
        return ''
    try:
        URLValidator(schemes=['https'])(url)
    except ValidationError:
        raise ValidationError(f'Row {row_number}: {field} must be a valid HTTPS URL.')
    if not url.lower().startswith('https://'):
        raise ValidationError(f'Row {row_number}: {field} must start with https://.')
    return url


def parse_resources(file):
    if file.size > 5 * 1024 * 1024:
        raise ValidationError('Excel file must not exceed 5 MB.')
    if not file.name.lower().endswith('.xlsx'):
        raise ValidationError('Upload an .xlsx Excel file.')
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError('The uploaded Excel file could not be read.') from exc
    sheet = workbook['Book Reference'] if 'Book Reference' in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or '').strip().lower() for value in next(rows, [])]
    missing = {'student resource link', 'teacher resource link'}.difference(headers)
    if 'isbn to paste' not in headers and 'isbn' not in headers:
        missing.add('isbn to paste')
    unknown = set(headers).difference(set(HEADERS + LEGACY_HEADERS))
    if missing:
        raise ValidationError(f'Missing columns: {", ".join(sorted(missing))}.')
    if unknown:
        raise ValidationError(f'Unknown columns: {", ".join(sorted(unknown))}. Use the current template.')
    isbn_header = 'isbn to paste' if 'isbn to paste' in headers else 'isbn'
    indexes = {header: headers.index(header) for header in [isbn_header, 'student resource link', 'teacher resource link']}
    parsed, errors, seen = [], [], set()
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, '') for value in values):
            continue
        try:
            isbn = normalize_isbn(values[indexes[isbn_header]] if indexes[isbn_header] < len(values) else '')
            if not isbn:
                raise ValidationError(f'Row {row_number}: enter a valid ISBN-10 or ISBN-13.')
            if isbn in seen:
                raise ValidationError(f'Row {row_number}: duplicate ISBN {isbn} in this file.')
            seen.add(isbn)
            product = Product.objects.filter(isbn=isbn).first()
            if not product:
                raise ValidationError(f'Row {row_number}: ISBN {isbn} does not exist in the books catalogue.')
            student_url = _validate_url(values[indexes['student resource link']], 'student resource link', row_number)
            teacher_url = _validate_url(values[indexes['teacher resource link']], 'teacher resource link', row_number)
            if not student_url and not teacher_url:
                raise ValidationError(f'Row {row_number}: enter at least one student or teacher resource link.')
            parsed.append({'product': product, 'student_url': student_url, 'teacher_url': teacher_url})
        except ValidationError as exc:
            errors.extend(exc.messages)
    if errors:
        raise ValidationError(errors[:100])
    if not parsed:
        raise ValidationError('The Book Reference sheet contains no rows.')
    return parsed


@transaction.atomic
def import_resources(file):
    rows = parse_resources(file)
    created = updated = 0
    for row in rows:
        _, was_created = DigitalResource.objects.update_or_create(
            product=row['product'],
            defaults={'student_url': row['student_url'], 'teacher_url': row['teacher_url']},
        )
        created += int(was_created)
        updated += int(not was_created)
    return {'created': created, 'updated': updated, 'total': len(rows)}
