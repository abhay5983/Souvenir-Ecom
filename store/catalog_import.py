from io import BytesIO
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.db import transaction
from django.utils.text import slugify
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from .models import Product


HEADERS = ['book', 'series', 'description', 'digital features', 'isbn', 'price', 'subject', 'cover photo link']
REQUIRED = {'book', 'series', 'isbn', 'price', 'subject'}


def template_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Books'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
    widths = [34, 28, 60, 34, 20, 14, 24, 50]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:H1'

    instructions = workbook.create_sheet('Instructions')
    instructions.append(['Column', 'Required', 'How to fill'])
    notes = {
        'book': 'YES|Individual book title, for example English Reader Class 5',
        'series': 'YES|Series name; use exactly the same spelling for all books in the series',
        'description': 'NO|Series description; repeat it on each row belonging to the series',
        'digital features': 'NO|Separate multiple features with |, for example Audio Book|Videos',
        'isbn': 'YES|Unique ISBN-10 or ISBN-13. Format this Excel cell as Text.',
        'price': 'YES|MRP in whole Indian rupees, without ₹ or commas',
        'subject': 'YES|For example English, Mathematics, Science or Hindi',
        'cover photo link': 'NO|Public HTTPS image URL; leave blank to show a generated cover',
    }
    for header in HEADERS:
        required, note = notes[header].split('|', 1)
        instructions.append([header, required, note])
    for cell in instructions[1]:
        cell.font = Font(bold=True)
    instructions.column_dimensions['A'].width = 24
    instructions.column_dimensions['B'].width = 12
    instructions.column_dimensions['C'].width = 90
    return workbook


def template_bytes():
    output = BytesIO()
    template_workbook().save(output)
    return output.getvalue()


def _text(value):
    return str(value or '').strip()


def _price(value, row_number):
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValidationError(f'Row {row_number}: price must be a whole number without symbols or commas.')
    if number < 1:
        raise ValidationError(f'Row {row_number}: price must be greater than zero.')
    return number


def _isbn(value, row_number):
    normalized = _text(value).replace('-', '').replace(' ', '').upper()
    if not ((len(normalized) == 10 and normalized[:9].isdigit() and (normalized[-1].isdigit() or normalized[-1] == 'X'))
            or (len(normalized) == 13 and normalized.isdigit())):
        raise ValidationError(f'Row {row_number}: ISBN must contain a valid 10 or 13 character ISBN format.')
    return normalized


def parse_catalogue(file):
    if file.size > 5 * 1024 * 1024:
        raise ValidationError('Excel file must not exceed 5 MB.')
    if not file.name.lower().endswith('.xlsx'):
        raise ValidationError('Upload an .xlsx Excel file.')
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError('The uploaded Excel file could not be read.') from exc
    sheet = workbook['Books'] if 'Books' in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    supplied_headers = [_text(value).lower() for value in next(rows, [])]
    missing = REQUIRED.difference(supplied_headers)
    if missing:
        raise ValidationError(f'Missing required columns: {", ".join(sorted(missing))}.')
    unknown = set(supplied_headers).difference(HEADERS)
    if unknown:
        raise ValidationError(f'Unknown columns: {", ".join(sorted(unknown))}. Use the current template.')
    indexes = {header: supplied_headers.index(header) for header in supplied_headers}
    parsed, errors, seen_isbns = [], [], set()
    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, '') for value in values):
            continue
        data = {header: values[index] if index < len(values) else None for header, index in indexes.items()}
        try:
            for field in REQUIRED:
                if data.get(field) in (None, ''):
                    raise ValidationError(f'Row {row_number}: {field} is required.')
            isbn = _isbn(data['isbn'], row_number)
            if isbn in seen_isbns:
                raise ValidationError(f'Row {row_number}: duplicate ISBN {isbn} in this file.')
            seen_isbns.add(isbn)
            image_url = _text(data.get('cover photo link'))
            if image_url:
                try:
                    URLValidator(schemes=['https'])(image_url)
                except ValidationError:
                    raise ValidationError(f'Row {row_number}: cover photo link must be a valid public HTTPS URL.')
            series_title = _text(data['series'])
            series_code = slugify(series_title)
            if not series_code:
                raise ValidationError(f'Row {row_number}: series must contain letters or numbers.')
            parsed.append({
                'sku': isbn,
                'isbn': isbn,
                'title': _text(data['book']),
                'series_code': series_code,
                'series_title': series_title,
                'series_slug': series_code,
                'imprint': 'Souvenir',
                'subject': _text(data['subject']),
                'category': 'Books',
                'description': _text(data.get('description')),
                'digital_features': [item.strip() for item in _text(data.get('digital features')).replace(',', '|').split('|') if item.strip()],
                'price_inr': _price(data['price'], row_number),
                'stock_quantity': 100,
                'hsn_code': '4901',
                'weight_kg': '0.300',
                'length_cm': 24,
                'breadth_cm': 18,
                'height_cm': 1,
                'cover_image_url': image_url,
                'featured': False,
                'public_visibility': True,
                'active': True,
            })
        except ValidationError as exc:
            errors.extend(exc.messages)
    if errors:
        raise ValidationError(errors[:100])
    if not parsed:
        raise ValidationError('The Books sheet contains no product rows.')
    return parsed


@transaction.atomic
def import_catalogue(file, replace_catalogue=False):
    rows = parse_catalogue(file)
    created = updated = 0
    imported_skus = []
    for row in rows:
        sku = row.pop('sku')
        _, was_created = Product.objects.update_or_create(sku=sku, defaults=row)
        imported_skus.append(sku)
        created += int(was_created)
        updated += int(not was_created)
    deactivated = 0
    if replace_catalogue:
        deactivated = Product.objects.exclude(sku__in=imported_skus).filter(active=True).update(active=False)
    return {'created': created, 'updated': updated, 'deactivated': deactivated, 'total': len(rows)}
