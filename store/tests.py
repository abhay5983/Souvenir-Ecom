import json
from io import BytesIO
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from .catalog_import import HEADERS, import_catalogue, template_bytes
from .models import DigitalLearningRequest, DigitalResource, GuestOrder, Product, PublicOutreachRequest
from .resource_import import HEADERS as RESOURCE_HEADERS, import_resources, template_bytes as resource_template_bytes


class GuestOrderApiTests(TestCase):
    def setUp(self):
        Product.objects.create(sku='book-1', title='Test Book', price_inr=300, stock_quantity=10)

    def post(self, path, body):
        return self.client.post(path, json.dumps(body), content_type='application/json')

    def test_guest_can_create_and_track_order(self):
        response = self.post('/api/orders', {
            'customer_name': 'Test Buyer', 'mobile': '9876543210', 'email': 'buyer@example.com',
            'address': {'line1': '12 Test Road', 'city': 'Delhi', 'state': 'Delhi', 'postal_code': '110001'},
            'lines': [{'product_id': 'book-1', 'quantity': 2}],
        })
        self.assertEqual(response.status_code, 201, response.content)
        created = response.json()
        self.assertEqual(created['subtotal'], 600)
        self.assertEqual(created['shippingCharge'], 200)
        self.assertEqual(created['deliveryType'], 'STANDARD')
        tracked = self.post('/api/orders/track', {'order_number': created['orderNumber'], 'mobile': '9876543210'})
        self.assertEqual(tracked.status_code, 200)
        self.assertEqual(tracked.json()['items'][0]['title'], 'Test Book')

    def test_frontend_price_is_not_accepted(self):
        response = self.post('/api/orders', {
            'customer_name': 'Test Buyer', 'mobile': '9876543210', 'email': 'buyer@example.com',
            'address': {'line1': '12 Test Road', 'city': 'Delhi', 'state': 'Delhi', 'postal_code': '110001'},
            'lines': [{'product_id': 'book-1', 'quantity': 1, 'unit_price': 1}],
        })
        self.assertEqual(response.json()['subtotal'], 300)

    def test_shipping_quote_uses_safe_fallback_before_credentials_exist(self):
        response = self.post('/api/shipping/serviceability', {
            'postal_code': '110001',
            'lines': [{'product_id': 'book-1', 'quantity': 1}],
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()['configured'])
        self.assertEqual(response.json()['shippingCharge'], 200)

    def test_priority_delivery_charge_is_calculated_by_backend(self):
        response = self.post('/api/orders', {
            'customer_name': 'Test Buyer', 'mobile': '9876543210', 'email': 'buyer@example.com',
            'address': {'line1': '12 Test Road', 'city': 'Delhi', 'state': 'Delhi', 'postal_code': '110001'},
            'delivery_type': 'PRIORITY', 'shipping_charge': 1,
            'lines': [{'product_id': 'book-1', 'quantity': 1}],
        })
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['shippingCharge'], 500)
        self.assertEqual(response.json()['total'], 800)
        self.assertEqual(response.json()['deliveryTimeline'], '2–3 working days')

    def test_unknown_delivery_option_is_rejected(self):
        response = self.post('/api/orders', {
            'customer_name': 'Test Buyer', 'mobile': '9876543210', 'email': 'buyer@example.com',
            'address': {'line1': '12 Test Road', 'city': 'Delhi', 'state': 'Delhi', 'postal_code': '110001'},
            'delivery_type': 'FREE',
            'lines': [{'product_id': 'book-1', 'quantity': 1}],
        })
        self.assertEqual(response.status_code, 400)

    def test_checkout_idempotency_does_not_create_duplicate_orders(self):
        body = {
            'idempotency_key': 'browser-attempt-123',
            'customer_name': 'Test Buyer', 'mobile': '9876543210', 'email': 'buyer@example.com',
            'address': {'line1': '12 Test Road', 'city': 'Delhi', 'state': 'Delhi', 'postal_code': '110001'},
            'lines': [{'product_id': 'book-1', 'quantity': 1}],
        }
        first = self.post('/api/orders', body)
        second = self.post('/api/orders', body)
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(first.json()['orderNumber'], second.json()['orderNumber'])
        self.assertEqual(GuestOrder.objects.filter(checkout_idempotency_key='browser-attempt-123').count(), 1)


class CatalogueImportTests(TestCase):
    def workbook_upload(self, values):
        workbook = load_workbook(BytesIO(template_bytes()))
        workbook['Books'].append(values)
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile('books.xlsx', output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_excel_import_creates_product_and_public_catalogue(self):
        row = {header: '' for header in HEADERS}
        row.update({
            'book': 'Real Book 1', 'series': 'Real Series', 'isbn': '9781234567890',
            'subject': 'English', 'description': 'A real series.',
            'digital features': 'Audio Book|Videos', 'price': 399,
            'cover photo link': 'https://cdn.example.com/real-book-1.jpg',
        })
        result = import_catalogue(self.workbook_upload([row[header] for header in HEADERS]))
        self.assertEqual(result['created'], 1)
        response = self.client.get('/api/catalogue')
        self.assertEqual(response.status_code, 200)
        series = response.json()['catalogue'][0]
        self.assertEqual(series['title'], 'Real Series')
        self.assertEqual(series['variants'][0]['id'], '9781234567890')
        self.assertEqual(series['variants'][0]['coverImageUrl'], 'https://cdn.example.com/real-book-1.jpg')
        self.assertEqual(series['digitalFeatures'], ['Audio Book', 'Videos'])


class DigitalResourceImportTests(TestCase):
    def setUp(self):
        self.product = Product.objects.create(
            sku='9781234567890', isbn='9781234567890', title='Resource Book',
            series_code='resource-series', series_slug='resource-series', series_title='Resource Series',
            subject='Science', price_inr=450,
        )

    def upload(self, isbn='9781234567890'):
        workbook = load_workbook(BytesIO(resource_template_bytes()))
        row = {
            'upload row': 2,
            'book name': 'Resource Book',
            'book code': 'RB-001',
            'isbn to paste': isbn,
            'student resource link': 'https://learn.example.com/student/book',
            'teacher resource link': 'https://learn.example.com/teacher/book',
        }
        headers = [cell.value for cell in workbook['Book Reference'][1]]
        workbook['Book Reference'].append([row.get(header, '') for header in headers])
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile('resources.xlsx', output.getvalue())

    def test_resource_upload_maps_existing_book_by_isbn(self):
        result = import_resources(self.upload())
        self.assertEqual(result['created'], 1)
        resource = DigitalResource.objects.get(product=self.product)
        self.assertEqual(resource.student_url, 'https://learn.example.com/student/book')

    def test_student_endpoint_does_not_expose_teacher_link(self):
        import_resources(self.upload())
        response = self.client.get('/api/digital-resources/student')
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['books'][0]['isbn'], self.product.isbn)
        self.assertNotIn('resourceUrl', payload['books'][0])
        self.assertNotIn('teacherUrl', payload['books'][0])
        self.assertNotIn('/teacher/', json.dumps(payload))

    def test_student_must_enter_matching_isbn_to_open_resource(self):
        import_resources(self.upload())
        denied = self.client.post('/api/digital-resources/student/access', json.dumps({
            'book_isbn': self.product.isbn, 'isbn': '9780000000000',
        }), content_type='application/json')
        self.assertEqual(denied.status_code, 403)
        allowed = self.client.post('/api/digital-resources/student/access', json.dumps({
            'book_isbn': self.product.isbn, 'isbn': '978-1-23456-789-0',
        }), content_type='application/json')
        self.assertEqual(allowed.status_code, 200)
        self.assertEqual(allowed.json()['resourceUrl'], 'https://learn.example.com/student/book')

    def test_teacher_lookup_returns_only_matching_link(self):
        import_resources(self.upload())
        response = self.client.post(
            '/api/digital-resources/teacher', json.dumps({
                'isbn': '978-1-23456-789-0', 'password': '9781234567890',
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['resourceUrl'], 'https://learn.example.com/teacher/book')

    def test_teacher_password_must_match_isbn(self):
        import_resources(self.upload())
        response = self.client.post('/api/digital-resources/teacher', json.dumps({
            'isbn': self.product.isbn, 'password': '9780000000000',
        }), content_type='application/json')
        self.assertEqual(response.status_code, 403)
        self.assertNotIn('resourceUrl', response.json())

    def test_teacher_bypass_returns_teacher_resource_catalogue(self):
        import_resources(self.upload())
        response = self.client.post('/api/digital-resources/teacher', json.dumps({
            'isbn': '9205406902', 'password': '9205406902',
        }), content_type='application/json')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['accessMode'], 'CATALOGUE')
        self.assertEqual(response.json()['books'][0]['title'], self.product.title)
        self.assertEqual(response.json()['books'][0]['resourceUrl'], 'https://learn.example.com/teacher/book')

    def test_teacher_bypass_requires_code_in_both_fields(self):
        response = self.client.post('/api/digital-resources/teacher', json.dumps({
            'isbn': '9205406902', 'password': self.product.isbn,
        }), content_type='application/json')
        self.assertEqual(response.status_code, 403)

    def test_unknown_isbn_does_not_create_resource(self):
        with self.assertRaisesMessage(Exception, 'does not exist in the books catalogue'):
            import_resources(self.upload('9780000000000'))
        self.assertEqual(DigitalResource.objects.count(), 0)


class DigitalLearningRequestTests(TestCase):
    def setUp(self):
        self.product = Product.objects.create(
            sku='9781234567890', isbn='9781234567890', title='Class 5 Science',
            series_code='science-world', series_slug='science-world', series_title='Science World',
            subject='Science', price_inr=450, public_visibility=True, active=True,
        )

    def post(self, overrides=None):
        payload = {
            'requesterRole': 'TEACHER_SCHOOL_ADMIN', 'requesterName': 'Test Teacher',
            'requesterDesignation': 'Teacher', 'schoolInstitutionName': 'Test School',
            'schoolBoard': 'CBSE', 'email': 'teacher@example.com', 'mobile': '9876543210',
            'stateCode': 'DL', 'pinCode': '110001', 'seriesId': 'science-world',
            'classIds': ['9781234567890'], 'resourceCodes': ['EBOOK', 'VIDEO'],
            'purpose': 'Classroom teaching',
            'usageDetails': 'The resources will support classroom teaching and revision.',
            'authorisedConfirmed': True, 'privacyAcknowledged': True,
        }
        payload.update(overrides or {})
        return self.client.post('/api/digital-resource-requests', json.dumps(payload), content_type='application/json')

    def test_public_request_is_saved_against_live_catalogue_book(self):
        response = self.post()
        self.assertEqual(response.status_code, 201, response.content)
        saved = DigitalLearningRequest.objects.get()
        self.assertEqual(saved.series_code, 'science-world')
        self.assertEqual(saved.selected_books[0]['isbn'], '9781234567890')
        self.assertTrue(response.json()['reference'].startswith('DLR-'))

    def test_request_rejects_book_from_another_series(self):
        response = self.post({'seriesId': 'wrong-series'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(DigitalLearningRequest.objects.count(), 0)

    def test_student_cannot_request_teacher_resource(self):
        response = self.post({
            'requesterRole': 'STUDENT', 'guardianContact': '9876543210',
            'resourceCodes': ['ANSWER_KEY'], 'email': '', 'mobile': '',
        })
        self.assertEqual(response.status_code, 403)

    def test_student_can_submit_with_guardian_email_only(self):
        response = self.post({
            'requesterRole': 'STUDENT',
            'requesterName': 'Test Student',
            'guardianContact': 'parent@example.com',
            'parentGuardianName': 'Test Parent',
            'guardianConsentConfirmed': True,
            'resourceCodes': ['EBOOK'],
            'email': '',
            'mobile': '',
        })
        self.assertEqual(response.status_code, 201, response.content)
        saved = DigitalLearningRequest.objects.get()
        self.assertEqual(saved.email, 'parent@example.com')
        self.assertEqual(saved.mobile, '')


class RequestHubTests(TestCase):
    def test_public_outreach_submission_is_persisted(self):
        response = self.client.post('/api/outreach-requests', json.dumps({
            'formType': 'GENERAL_ENQUIRY',
            'values': {'fullName': 'Public Customer', 'email': 'customer@example.com', 'mobile': '9876543210',
                       'city': 'Delhi', 'state': 'Delhi', 'pincode': '110001', 'topic': 'Books and catalogue',
                       'description': 'Please help with my catalogue question.'},
            'files': [], 'honeypot': '',
        }), content_type='application/json')
        self.assertEqual(response.status_code, 201, response.content)
        saved = PublicOutreachRequest.objects.get()
        self.assertEqual(saved.requester_name, 'Public Customer')
        self.assertTrue(response.json()['publicReference'].startswith('OUT-'))

    def test_public_user_can_view_combined_request_hub(self):
        PublicOutreachRequest.objects.create(form_type='GENERAL_ENQUIRY', form_title='General Enquiry', category='General Support')
        response = self.client.get('/api/request-hub')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['requests'][0]['source'], 'OUTREACH')

    def test_public_user_can_update_request_from_dashboard(self):
        item = PublicOutreachRequest.objects.create(
            form_type='GENERAL_ENQUIRY', form_title='General Enquiry', category='General Support',
        )
        response = self.client.post(
            f'/api/request-hub/OUTREACH/{item.id}',
            json.dumps({'status': 'IN_REVIEW', 'assignedTeam': 'Customer Care', 'note': 'Called customer.'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200, response.content)
        item.refresh_from_db()
        self.assertEqual(item.status, 'IN_REVIEW')
        self.assertEqual(item.internal_notes[0]['actor'], 'Request dashboard')

    def test_customer_order_appears_and_can_be_updated_in_request_hub(self):
        product = Product.objects.create(sku='hub-book', title='Hub Book', price_inr=250)
        order = GuestOrder.objects.create(
            customer_name='Order Customer', mobile='9876543210', email='order@example.com',
            address={'line1': '1 Book Road', 'city': 'Delhi', 'state': 'Delhi', 'postal_code': '110001'},
            subtotal=250, shipping_charge=200, total=450,
        )
        order.items.create(product=product, sku=product.sku, title=product.title, unit_price=250, quantity=1, line_total=250)
        response = self.client.get('/api/request-hub')
        saved = next(item for item in response.json()['requests'] if item['source'] == 'ORDER')
        self.assertEqual(saved['reference'], order.order_number)
        self.assertEqual(saved['payload']['items'][0]['title'], 'Hub Book')
        updated = self.client.post(
            f'/api/request-hub/ORDER/{order.id}', json.dumps({'status': 'PROCESSING'}),
            content_type='application/json',
        )
        self.assertEqual(updated.status_code, 200, updated.content)
        self.assertEqual(updated.json()['status'], 'PROCESSING')
